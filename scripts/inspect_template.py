from pptx import Presentation

path = r"c:\Users\Neel Parikh\Downloads\KJSSE  PPT Template _From Jan2025_FinalAprroved (1).pptx"
prs = Presentation(path)
print("slide_width", prs.slide_width)
print("slide_height", prs.slide_height)
print("slide_count", len(prs.slides))
print("layouts", len(prs.slide_layouts))
for i, layout in enumerate(prs.slide_layouts):
    print(f"\nLAYOUT {i}: name={layout.name}")
    for j, ph in enumerate(layout.placeholders):
        print(
            f"  ph[{j}] idx={ph.placeholder_format.idx} type={ph.placeholder_format.type} "
            f"name={ph.name} left={ph.left} top={ph.top} w={ph.width} h={ph.height}"
        )

print("\n=== EXISTING SLIDES ===")
for si, slide in enumerate(prs.slides):
    print(f"\nSLIDE {si+1} layout={slide.slide_layout.name}")
    for shape in slide.shapes:
        info = f"  shape={shape.shape_type} name={shape.name}"
        if shape.has_text_frame:
            texts = []
            for p in shape.text_frame.paragraphs:
                t = "".join([r.text for r in p.runs]) or p.text
                if t.strip():
                    texts.append(t.strip()[:160])
            if texts:
                info += " | " + " || ".join(texts)
        print(info)
