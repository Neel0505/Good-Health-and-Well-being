"""
Build FitTrack WPL Mini Project PPT using the official KJSSE/SVU template.
"""
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import shutil
import os

TEMPLATE = r"c:\Users\Neel Parikh\Downloads\KJSSE  PPT Template _From Jan2025_FinalAprroved (1).pptx"
OUT_DIR = r"c:\FitTrack\docs"
OUT_PPTX = os.path.join(OUT_DIR, "FitTrack_WPL_Mini_Project_Presentation.pptx")
OUT_PDF = os.path.join(OUT_DIR, "FitTrack_WPL_Mini_Project_Presentation.pdf")

# Somaiya-aligned accent (deep maroon/orange typical of SVU branding)
SOMAIYA_MAROON = RGBColor(0x8B, 0x1E, 0x3F)
SOMAIYA_ORANGE = RGBColor(0xE8, 0x6C, 0x00)
DARK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x0F, 0x76, 0x6E)


def delete_all_slides(prs):
    """Remove every existing slide while preserving slide master/theme."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def set_run_font(run, size=18, bold=False, color=DARK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def clear_placeholder(ph):
    tf = ph.text_frame
    tf.clear()
    return tf


def set_title(slide, text, size=28, align=PP_ALIGN.LEFT):
    ph = slide.shapes.title
    tf = clear_placeholder(ph)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=True, color=SOMAIYA_MAROON)


def add_bullets(text_frame, items, size=16, level0_bold=False, color=DARK):
    """items: list of str or (text, level)"""
    # Ensure first paragraph exists empty
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        if first:
            p = text_frame.paragraphs[0]
            first = False
        else:
            p = text_frame.add_paragraph()
        p.level = level
        p.space_after = Pt(6)
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = text
        set_run_font(run, size=size, bold=(level0_bold and level == 0), color=color)


def set_footer(slide, footer_text="FitTrack | WPL Mini Project | Batch A-3"):
    for shape in slide.placeholders:
        try:
            idx = shape.placeholder_format.idx
        except Exception:
            continue
        if idx == 11:  # footer
            tf = clear_placeholder(shape)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = footer_text
            set_run_font(run, size=10, color=MUTED)
        elif idx == 10:  # date
            tf = clear_placeholder(shape)
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "AY 2025-26"
            set_run_font(run, size=10, color=MUTED)


def add_textbox(slide, left, top, width, height, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    return box


def add_card(slide, left, top, width, height, title, lines, fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or RGBColor(0xF7, 0xF2, 0xF4)
    shape.line.color.rgb = RGBColor(0xE5, 0xD5, 0xDB)
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(8)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run_font(run, size=14, bold=True, color=SOMAIYA_MAROON)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = line
        set_run_font(run, size=12, color=DARK)
    return shape


def body_ph(slide, idx=1):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == idx:
            return shape
    # fallback
    return slide.placeholders[1]


def build():
    os.makedirs(OUT_DIR, exist_ok=True)
    shutil.copy2(TEMPLATE, OUT_PPTX)
    prs = Presentation(OUT_PPTX)
    delete_all_slides(prs)

    L_TITLE = prs.slide_layouts[0]
    L_OBJECT = prs.slide_layouts[1]
    L_TWO = prs.slide_layouts[4]
    L_SECTION = prs.slide_layouts[5]
    L_TITLE_ONLY = prs.slide_layouts[7]
    L_CAPTION = prs.slide_layouts[2]

    # ========== 1. Title + Team ==========
    s = prs.slides.add_slide(L_TITLE)
    set_title(s, "FitTrack: Personal Health, Nutrition\nand Fitness Tracker", size=30, align=PP_ALIGN.CENTER)
    # subtitle placeholder
    sub = None
    for shape in s.placeholders:
        if shape.placeholder_format.idx == 1:
            sub = shape
            break
    tf = clear_placeholder(sub)
    lines = [
        "Web Programming Laboratory (WPL) Mini Project",
        "Theme: SDG 3 - Good Health and Well-Being",
        "",
        "Neel Parikh  |  Roll No. 16014024052",
        "Shourya Ranjan  |  Roll No. 16014024057",
        "",
        "Batch A-3  |  Branch: EXCP",
        "K. J. Somaiya College of Engineering, Mumbai",
    ]
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        bold = "Neel" in line or "Shourya" in line or line.startswith("FitTrack") or "WPL" in line
        set_run_font(run, size=14 if line else 8, bold=bold, color=DARK)
    set_footer(s)

    # ========== 2. Agenda ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Agenda")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "1. Project Title and Team Details",
            "2. Problem Statement",
            "3. System Architecture",
            "4. Features of the Designed System",
            "5. Results and Demonstration Outcomes",
            "6. Conclusion",
            "7. References",
        ],
        size=18,
    )
    set_footer(s)

    # ========== 3. Problem Statement ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Problem Statement")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "Many individuals struggle to consistently track nutrition, workouts and daily health habits in one place.",
            "Existing tools are often fragmented, complex, or lack an educational, student-friendly implementation.",
            "There is a need for a unified web application that helps users:",
            ("Monitor calories, protein and meals", 1),
            ("Record strength and cardio workouts", 1),
            ("Track water, sleep and weight", 1),
            ("Visualize daily progress through dashboards and charts", 1),
            "FitTrack addresses this need as a complete MERN-stack solution aligned with SDG 3: Good Health and Well-Being.",
            "Note: FitTrack is an educational tracking tool and does not provide medical diagnosis.",
        ],
        size=15,
    )
    set_footer(s)

    # ========== 4. Objectives ==========
    s = prs.slides.add_slide(L_TWO)
    set_title(s, "Project Objectives & SDG Alignment")
    left = clear_placeholder(body_ph(s, 1))
    add_bullets(
        left,
        [
            "Objectives",
            ("Build a secure user authentication system using JWT", 1),
            ("Enable nutrition search, logging and daily targets", 1),
            ("Support workout creation with exercise library", 1),
            ("Provide health habit tracking and progress charts", 1),
            ("Demonstrate modular MERN architecture for viva", 1),
        ],
        size=14,
        level0_bold=True,
    )
    right = clear_placeholder(body_ph(s, 2))
    add_bullets(
        right,
        [
            "SDG 3 Contribution",
            ("Promote awareness of daily nutrition choices", 1),
            ("Encourage consistent physical activity tracking", 1),
            ("Support hydration, sleep and weight monitoring", 1),
            ("Help users build healthier everyday habits", 1),
            ("Educational focus - estimates, not medical advice", 1),
        ],
        size=14,
        level0_bold=True,
    )
    set_footer(s)

    # ========== 5. System Architecture ==========
    s = prs.slides.add_slide(L_TITLE_ONLY)
    set_title(s, "System Architecture (MERN Stack)")
    # Architecture flow cards
    cards = [
        (Inches(0.7), "React Client\n(Vite + Bootstrap)", "UI, Forms, Charts\nProtected Routes"),
        (Inches(3.55), "REST API\n(Express.js)", "JWT Auth, Validation\nControllers & Routes"),
        (Inches(6.4), "Data Layer\n(Mongoose)", "Schemas, Indexes\nCRUD Operations"),
        (Inches(9.25), "MongoDB\nDatabase", "Users, Foods, Logs\nWorkouts, Health"),
    ]
    top = Inches(1.7)
    for left, title, body in cards:
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.55), Inches(2.35))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xF8, 0xF1, 0xF4)
        shape.line.color.rgb = SOMAIYA_MAROON
        shape.line.width = Pt(1.5)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(8)
        tf.margin_right = Pt(8)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        set_run_font(run, size=13, bold=True, color=SOMAIYA_MAROON)
        for line in body.split("\n"):
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_before = Pt(6)
            run = p.add_run()
            run.text = line
            set_run_font(run, size=11, color=DARK)

    # arrows between cards using chevrons as text
    for x in [Inches(3.2), Inches(6.05), Inches(8.9)]:
        add_textbox(s, x, Inches(2.5), Inches(0.4), Inches(0.4), ">", size=22, bold=True, color=SOMAIYA_ORANGE, align=PP_ALIGN.CENTER)

    add_textbox(
        s,
        Inches(0.7),
        Inches(4.25),
        Inches(12),
        Inches(0.9),
        "Data Flow: Browser  >  Axios REST Calls  >  Express Middleware (JWT)  >  Mongoose Models  >  MongoDB Collections",
        size=13,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    set_footer(s)

    # ========== 6. Architecture detail ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Modular Backend & Frontend Design")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "Backend Modules (completely modular with clear function separation)",
            ("Models: User, Food, FoodLog, Exercise, Workout, HealthLog", 1),
            ("Controllers: Auth, Users, Foods, FoodLogs, Workouts, Health, Dashboard, Progress", 1),
            ("Middleware: JWT protect, express-validator, centralized error handler", 1),
            ("Utilities: nutrition target calculator (Mifflin-St Jeor), token generation", 1),
            "Frontend Modules",
            ("Pages, reusable components, AuthContext, Axios service layer, layouts", 1),
            ("Protected routes for dashboard pages; public routes for Home/About/Auth", 1),
            "This modular structure makes the project viva-friendly and maintainable.",
        ],
        size=14,
        level0_bold=True,
    )
    set_footer(s)

    # ========== 7. Technology Stack ==========
    s = prs.slides.add_slide(L_TWO)
    set_title(s, "Technology Stack")
    left = clear_placeholder(body_ph(s, 1))
    add_bullets(
        left,
        [
            "Frontend",
            ("React + Vite (JavaScript)", 1),
            ("Bootstrap 5 + Bootstrap Icons", 1),
            ("React Router, Axios", 1),
            ("Chart.js / react-chartjs-2", 1),
            ("Responsive custom CSS theme", 1),
        ],
        size=15,
        level0_bold=True,
    )
    right = clear_placeholder(body_ph(s, 2))
    add_bullets(
        right,
        [
            "Backend & Database",
            ("Node.js + Express.js REST API", 1),
            ("MongoDB + Mongoose ODM", 1),
            ("JWT Authentication", 1),
            ("bcryptjs password hashing", 1),
            ("express-validator + CORS + dotenv", 1),
        ],
        size=15,
        level0_bold=True,
    )
    set_footer(s)

    # ========== 8. Features overview ==========
    s = prs.slides.add_slide(L_TITLE_ONLY)
    set_title(s, "Features of the Designed System")
    features = [
        ("Pillar 1", "Nutrition &\nProtein Tracking", "Food DB, meal logs,\ncalorie & protein targets"),
        ("Pillar 2", "Workout &\nExercise Tracking", "Exercise library,\nsets/reps/cardio logs"),
        ("Pillar 3", "Health Dashboard\n& Progress", "Water, sleep, weight,\nChart.js analytics"),
    ]
    for i, (tag, title, desc) in enumerate(features):
        left = Inches(0.8 + i * 4.0)
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.75), Inches(3.6), Inches(3.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = SOMAIYA_MAROON
        shape.line.width = Pt(1.75)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.margin_right = Pt(12)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = tag
        set_run_font(run, size=12, bold=True, color=SOMAIYA_ORANGE)
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(10)
        run = p.add_run()
        run.text = title
        set_run_font(run, size=16, bold=True, color=SOMAIYA_MAROON)
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(14)
        run = p.add_run()
        run.text = desc
        set_run_font(run, size=13, color=DARK)
    set_footer(s)

    # ========== 9. Nutrition features ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Feature Deep Dive - Nutrition Module")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "Seeded food database with categories: Protein, Grains, Fruits, Vegetables, Snacks",
            "Search + category filters; select food, quantity and meal type",
            "Dynamic nutrition scaling (example: Chicken Breast 150 g -> 247.5 kcal, 46.5 g protein)",
            "Food logs grouped by Breakfast, Lunch, Dinner and Snacks",
            "Custom food creation with validation",
            "Daily calorie and protein targets using Mifflin-St Jeor BMR x activity x goal",
            "Progress bars comparing intake vs target on Nutrition page and Dashboard",
        ],
        size=15,
    )
    set_footer(s)

    # ========== 10. Workout + Health ==========
    s = prs.slides.add_slide(L_TWO)
    set_title(s, "Feature Deep Dive - Workouts & Health")
    left = clear_placeholder(body_ph(s, 1))
    add_bullets(
        left,
        [
            "Workout Module",
            ("Exercise library with search & filters", 1),
            ("Categories: Strength, Cardio, Bodyweight, Flexibility", 1),
            ("Log sets, reps, weight or duration/distance", 1),
            ("Save completed workouts and view history", 1),
            ("Delete and review workout details", 1),
        ],
        size=14,
        level0_bold=True,
    )
    right = clear_placeholder(body_ph(s, 2))
    add_bullets(
        right,
        [
            "Health & Progress",
            ("Daily water intake (ml)", 1),
            ("Sleep duration (hours)", 1),
            ("Body weight logging", 1),
            ("Dashboard summary cards", 1),
            ("Charts for 7 / 30 / 90 days (real data only)", 1),
        ],
        size=14,
        level0_bold=True,
    )
    set_footer(s)

    # ========== 11. Auth & Security ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Authentication, Security & Validation")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "User registration and login with email uniqueness checks",
            "Passwords hashed using bcryptjs (never stored in plaintext)",
            "JWT tokens issued on login/register; Authorization: Bearer middleware",
            "Protected frontend routes (/app/*) and protected REST endpoints",
            "Users can access only their own food logs, workouts and health data",
            "Input validation on both frontend forms and backend (express-validator)",
            "Centralized API error responses: { success, message/data }",
            "Secrets and Mongo URI stored in environment variables (.env)",
        ],
        size=15,
    )
    set_footer(s)

    # ========== 12. Results ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Results - Working System Outcomes")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "Successfully developed a complete, runnable MERN application: FitTrack",
            "Verified end-to-end flows:",
            ("User registration, login and profile target calculation", 1),
            ("Food search, meal logging and daily nutrition totals", 1),
            ("Workout creation from exercise library", 1),
            ("Health habit logging and dashboard aggregation", 1),
            ("Progress charts generated from real MongoDB data", 1),
            "Seed data: 34 foods and 22 exercises for immediate usability",
            "Responsive UI with professional health/fitness visual design",
            "Application runs locally: Client :5173 | API :5000 | MongoDB :27017",
        ],
        size=14,
    )
    set_footer(s)

    # ========== 13. Test Cases (rubric excellence) ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Testing - Sample Test Cases")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "TC01 Register with valid details -> Account created, JWT returned",
            "TC02 Register with duplicate email -> 400 validation error",
            "TC03 Login with wrong password -> 401 Unauthorized",
            "TC04 Access /api/dashboard without token -> 401 blocked",
            "TC05 Log Chicken Breast 150 g -> Calories 247.5, Protein 46.5 computed correctly",
            "TC06 Create workout with sets/reps -> Saved and visible in history",
            "TC07 Update profile age/height/weight -> Calorie & protein targets recalculated",
            "TC08 Progress page with no prior data -> Empty state (no fabricated charts)",
            "TC09 Mobile/responsive layout -> Sidebar collapses; no horizontal scroll",
        ],
        size=13,
    )
    set_footer(s)

    # ========== 14. Conclusion ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "Conclusion")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "FitTrack successfully demonstrates a full-stack MERN solution for personal health tracking.",
            "The system integrates three core pillars - Nutrition, Workouts and Health Progress - into one cohesive application.",
            "Modular architecture, JWT security, validation and real MongoDB persistence make the project robust and viva-ready.",
            "Aligned with SDG 3 by promoting awareness of everyday well-being habits through practical digital tracking.",
            "Future scope: meal templates, workout plans, PDF reports and enhanced analytics.",
            "Limitation: nutritional calculations are educational estimates and not medical advice.",
        ],
        size=15,
    )
    set_footer(s)

    # ========== 15. References ==========
    s = prs.slides.add_slide(L_OBJECT)
    set_title(s, "References")
    tf = clear_placeholder(body_ph(s))
    add_bullets(
        tf,
        [
            "United Nations - Sustainable Development Goal 3: Good Health and Well-Being. https://sdgs.un.org/goals/goal3",
            "MongoDB Documentation. https://www.mongodb.com/docs/",
            "Express.js Documentation. https://expressjs.com/",
            "React Documentation. https://react.dev/",
            "Mongoose ODM Documentation. https://mongoosejs.com/docs/",
            "Chart.js Documentation. https://www.chartjs.org/docs/",
            "Mifflin, M. D. et al. - A new predictive equation for resting energy expenditure (Mifflin-St Jeor).",
            "Bootstrap 5 Documentation. https://getbootstrap.com/docs/5.3/",
            "JSON Web Token (RFC 7519) and bcrypt password hashing references.",
        ],
        size=13,
    )
    set_footer(s)

    # ========== 16. Thank You ==========
    s = prs.slides.add_slide(L_TITLE)
    set_title(s, "Thank You", size=40, align=PP_ALIGN.CENTER)
    sub = None
    for shape in s.placeholders:
        if shape.placeholder_format.idx == 1:
            sub = shape
            break
    tf = clear_placeholder(sub)
    thanks = [
        "Questions & Discussion Welcome",
        "",
        "Neel Parikh - 16014024052",
        "Shourya Ranjan - 16014024057",
        "Batch A-3 | EXCP | WPL Mini Project",
        "",
        "K. J. Somaiya College of Engineering",
        "Somaiya Vidyavihar University",
    ]
    first = True
    for line in thanks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        set_run_font(run, size=16 if line else 8, bold=("Neel" in line or "Shourya" in line or "Questions" in line), color=DARK)
    set_footer(s)

    prs.save(OUT_PPTX)
    print(f"Saved PPTX: {OUT_PPTX}")
    print(f"Slide count: {len(prs.slides)}")
    return OUT_PPTX, OUT_PDF


if __name__ == "__main__":
    build()
