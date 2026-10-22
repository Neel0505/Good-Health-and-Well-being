from pptx import Presentation

path = r"c:\FitTrack\docs\FitTrack_WPL_Mini_Project_Presentation.pptx"
prs = Presentation(path)
print("slides", len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = "\n".join(
                ["".join([r.text for r in p.runs]) for p in shape.text_frame.paragraphs]
            ).strip()
            if t:
                texts.append(t.replace("\n", " | ")[:140])
    print(f"\n#{i} [{slide.slide_layout.name}]")
    for t in texts[:6]:
        print(" ", t)
