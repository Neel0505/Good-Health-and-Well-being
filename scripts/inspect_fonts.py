from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

path = r"c:\Users\Neel Parikh\Downloads\KJSSE  PPT Template _From Jan2025_FinalAprroved (1).pptx"
prs = Presentation(path)

# Inspect first slide title placeholder font
slide = prs.slides[0]
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            font = r.font
            print(
                shape.name,
                repr(r.text[:40]),
                "name=", font.name,
                "size=", font.size,
                "bold=", font.bold,
                "color=", getattr(font.color, "rgb", None),
            )

print("\n--- OBJECT layout sample font from slide master ---")
layout = prs.slide_layouts[1]
for shape in layout.placeholders:
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                print(shape.name, repr(r.text[:40]), font.name if (font:=r.font) else None, font.size)
            if not p.runs and p.text:
                print("para text", p.text)
